"""Documentos no PC: upload, leitura, resumo, tópicos e Q&A sobre o conteúdo.

Formatos lidos: PDF, DOCX, TXT/Markdown, CSV/TSV, XLSX, HTML e JSON.
Outros (PPTX, ODT, EPUB…) retornam erro amigável pedindo conversão para PDF.

Tudo com imports opcionais: se a lib extratora não estiver instalada no PC,
o erro diz exatamente qual `pip install` fazer — nada quebra o agente.
"""
import csv
import hashlib
import html
import io
import json
import logging
import os
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

log = logging.getLogger("JULIA-docs")

MAX_BYTES = 15 * 1024 * 1024
MAX_TEXTO_DOC = 200_000
MAX_ZIP_ENTRIES = 1000
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 64 * 1024 * 1024
MAX_ZIP_RATIO = 200
TAM_CHUNK = 3000
OVERLAP = 300
STOP = frozenset(("que com uma para das dos por seu sua isto isso entre como mas foi ser tem são foi ela ele eles elas meu minha teu tua seu sua nosso nossa este esta esse essa aquele aquela muito mais menos sobre após depois onde quando qual quais todo toda todos todas também pode fazer faz parte vez vezes".split()))

EXT_OK = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".tsv",
          ".xlsx", ".html", ".htm", ".json"}
DOC_VERSION = 1


class DocErro(ValueError):
    def __init__(self, mensagem, codigo="erro"):
        super().__init__(mensagem)
        self.codigo = codigo


def agora_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def docs_dir() -> Path:
    base = os.environ.get("JULIA_DOCS_PATH", "")
    return Path(base) if base else (Path.home() / ".JULIA-mobile" / "docs")


DONO_PADRAO = "local"


def criar_tabelas(db) -> None:
    db.execute("""CREATE TABLE IF NOT EXISTS documents (
        id TEXT PRIMARY KEY, name TEXT NOT NULL, ext TEXT NOT NULL,
        size INTEGER NOT NULL, chars INTEGER NOT NULL, chunks INTEGER NOT NULL,
        summary TEXT, topics TEXT, conversation_id TEXT NOT NULL,
        created_at TEXT NOT NULL, owner_id TEXT NOT NULL DEFAULT 'local'
    )""")
    db.execute("""CREATE TABLE IF NOT EXISTS doc_chunks (
        doc_id TEXT NOT NULL, idx INTEGER NOT NULL, text TEXT NOT NULL,
        owner_id TEXT NOT NULL DEFAULT 'local',
        PRIMARY KEY (doc_id, idx)
    )""")
    _ensure_columns(db, "documents", {
        "owner_id": "TEXT NOT NULL DEFAULT 'local'",
        "title": "TEXT NOT NULL DEFAULT ''",
        "extracted_meta": "TEXT NOT NULL DEFAULT '{}'",
        "content_hash": "TEXT NOT NULL DEFAULT ''",
        "version": "INTEGER NOT NULL DEFAULT 1",
        "source": "TEXT NOT NULL DEFAULT ''",
    })
    _ensure_columns(db, "doc_chunks", {
        "owner_id": "TEXT NOT NULL DEFAULT 'local'",
        "chunk_id": "TEXT NOT NULL DEFAULT ''",
        "section": "TEXT NOT NULL DEFAULT ''",
        "page": "INTEGER",
        "content_hash": "TEXT NOT NULL DEFAULT ''",
        "version": "INTEGER NOT NULL DEFAULT 1",
        "source": "TEXT NOT NULL DEFAULT ''",
        "table_hint": "INTEGER NOT NULL DEFAULT 0",
        "overlap": "INTEGER NOT NULL DEFAULT 0",
    })
    _migrate_chunk_metadata(db)
    db.execute("CREATE INDEX IF NOT EXISTS idx_docs_conv ON documents(conversation_id, created_at)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_doc_chunks_owner ON doc_chunks(owner_id, doc_id, idx)")
    try:
        db.execute("""CREATE VIRTUAL TABLE IF NOT EXISTS doc_chunks_fts USING fts5(
            chunk_id UNINDEXED, doc_id UNINDEXED, owner_id UNINDEXED, idx UNINDEXED, text,
            tokenize='unicode61 remove_diacritics 2'
        )""")
        # Idempotently backfill old rows without rebuilding the index on every startup.
        db.execute("""INSERT INTO doc_chunks_fts(chunk_id, doc_id, owner_id, idx, text)
                      SELECT c.chunk_id, c.doc_id, c.owner_id, c.idx, c.text FROM doc_chunks c
                      WHERE NOT EXISTS (
                          SELECT 1 FROM doc_chunks_fts f
                          WHERE f.doc_id=c.doc_id AND f.idx=c.idx
                      )""")
    except Exception as exc:
        log.info("SQLite FTS5 indisponível; usando busca lexical: %s", exc)
    db.commit()


def _ensure_columns(db, table: str, columns: dict) -> None:
    existing = {row[1] for row in db.execute(f"PRAGMA table_info({table})")}
    for name, declaration in columns.items():
        if name not in existing:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {declaration}")


def _migrate_chunk_metadata(db) -> None:
    """Backfill provenance once for rows created by the pre-chunk schema."""
    docs = db.execute("SELECT id, name, owner_id, chars, extracted_meta, content_hash "
                      "FROM documents").fetchall()
    for doc_id, name, owner, chars, raw_meta, doc_hash in docs:
        rows = db.execute("""SELECT idx, text, chunk_id, content_hash, source
                             FROM doc_chunks WHERE doc_id=? AND owner_id=? ORDER BY idx""",
                          (doc_id, owner)).fetchall()
        if not rows:
            continue
        legacy = any(not row[2] for row in rows)
        assembled = []
        previous = ""
        previous_idx = None
        for idx, text, chunk_id, digest, source in rows:
            overlap = 0
            if legacy and previous_idx is not None and idx == previous_idx + 1:
                for length in range(min(OVERLAP, len(previous)), 0, -1):
                    if previous.endswith(text[:length]):
                        overlap = length
                        break
            full_hash = digest or hashlib.sha256(text.encode("utf-8")).hexdigest()
            stable_id = chunk_id or f"{doc_id}:{idx}"
            db.execute("""UPDATE doc_chunks SET chunk_id=?, content_hash=?, source=?,
                          overlap=CASE WHEN chunk_id='' THEN ? ELSE overlap END
                          WHERE doc_id=? AND idx=? AND owner_id=?""",
                       (stable_id, full_hash, source or name, overlap, doc_id, idx, owner))
            assembled.append(text[overlap:])
            previous, previous_idx = text, idx
        if legacy:
            normalized = "".join(assembled)
            meta = json.loads(raw_meta or "{}")
            if not meta:
                meta = {"sections": [], "pages": 0, "table_hints": [],
                        "truncated": chars >= MAX_TEXTO_DOC, "extraction": "legacy_unknown"}
            db.execute("""UPDATE documents SET title=CASE WHEN title='' THEN ? ELSE title END,
                          source=CASE WHEN source='' THEN ? ELSE source END,
                          content_hash=CASE WHEN content_hash='' THEN ? ELSE content_hash END,
                          extracted_meta=CASE WHEN extracted_meta='{}' THEN ? ELSE extracted_meta END
                          WHERE id=? AND owner_id=?""",
                       (Path(name).stem[:200], name,
                        doc_hash or hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
                        json.dumps(meta, ensure_ascii=False), doc_id, owner))


def _tokens(s: str) -> set:
    return {t for t in re.findall(r"[a-zà-ú0-9]{3,}", (s or "").lower()) if t not in STOP}


# ── Extração ──────────────────────────────────────────────────────────
def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, ValueError):
            continue
    return raw.decode("utf-8", "ignore")


def _validar_zip_ooxml(raw: bytes, ext: str) -> None:
    """Check package structure and expansion limits before any OOXML parser runs."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise DocErro("Arquivo compactado contém entradas demais.", "zip_limite")
            names = set()
            total = 0
            for info in entries:
                name = info.filename.replace("\\", "/")
                parts = name.split("/")
                if (name.startswith("/") or re.match(r"^[a-zA-Z]:", name)
                        or any(part in (".", "..") for part in parts)):
                    raise DocErro("Caminho interno inválido no arquivo compactado.", "zip_caminho")
                if info.flag_bits & 0x1:
                    raise DocErro("Arquivo compactado criptografado não é aceito.", "zip_criptografado")
                if info.file_size > MAX_ZIP_MEMBER_BYTES:
                    raise DocErro("Parte compactada excede o limite de 16 MiB.", "zip_limite")
                if total + info.file_size > MAX_ZIP_TOTAL_BYTES:
                    raise DocErro("Conteúdo descompactado excede o limite de 64 MiB.", "zip_limite")
                if info.file_size and (not info.compress_size
                        or info.file_size / info.compress_size > MAX_ZIP_RATIO):
                    raise DocErro("Taxa de compactação suspeita; arquivo rejeitado.", "zip_ratio")
                if name.endswith("/"):
                    continue
                if name in names:
                    raise DocErro("Nomes de partes duplicados no arquivo compactado.", "zip_invalido")
                names.add(name)
                # Count bytes actually decompressed, not only attacker-controlled ZIP headers.
                content = bytearray() if name.lower().endswith((".xml", ".rels")) else None
                member_size = 0
                with archive.open(info) as member:
                    while True:
                        block = member.read(min(64 * 1024, MAX_ZIP_MEMBER_BYTES - member_size + 1))
                        if not block:
                            break
                        member_size += len(block)
                        total += len(block)
                        if member_size > MAX_ZIP_MEMBER_BYTES or total > MAX_ZIP_TOTAL_BYTES:
                            raise DocErro("Conteúdo descompactado excede os limites permitidos.", "zip_limite")
                        if content is not None:
                            content.extend(block)
                if content is not None:
                    if re.search(rb"<!\s*(?:DOCTYPE|ENTITY)\b", content, re.I):
                        raise DocErro("DTD/entidades XML não são permitidos.", "xml_unsafe")
                    if name.lower().endswith(".rels") and re.search(
                            rb"\bTargetMode\s*=\s*['\"]External['\"]", content, re.I):
                        raise DocErro("Referências externas em XML não são permitidas.", "xml_unsafe")
            required = ({"[Content_Types].xml", "word/document.xml"} if ext == ".docx"
                        else {"[Content_Types].xml", "xl/workbook.xml"})
            if not required.issubset(names):
                raise DocErro("Arquivo OOXML inválido para a extensão informada.", "conteudo_invalido")
    except DocErro:
        raise
    except (zipfile.BadZipFile, OSError, RuntimeError, ValueError):
        raise DocErro("Arquivo compactado inválido ou corrompido.", "conteudo_invalido") from None


def _validar_conteudo(ext: str, raw: bytes) -> None:
    if ext == ".pdf":
        if not raw.lstrip().startswith(b"%PDF-"):
            raise DocErro("O conteúdo não corresponde a um PDF.", "conteudo_invalido")
    elif ext in (".docx", ".xlsx"):
        if not raw.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
            raise DocErro("O conteúdo não corresponde a um pacote Office.", "conteudo_invalido")
        _validar_zip_ooxml(raw, ext)
    elif ext in (".txt", ".md", ".markdown", ".csv", ".tsv", ".html", ".htm", ".json"):
        controls = sum(byte < 32 and byte not in (9, 10, 12, 13) or byte == 127
                       for byte in raw)
        if b"\x00" in raw or controls > max(1, len(raw) // 100):
            raise DocErro("O conteúdo contém bytes binários incompatíveis com texto.", "conteudo_invalido")


def _extrair_pdf(raw: bytes) -> str:
    partes = []
    titulo = ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        title = getattr(reader.metadata, "title", None) if reader.metadata else None
        if title:
            titulo = f"# {str(title)[:200]}\n"
        for i, pg in enumerate(reader.pages[:200]):
            try:
                t = (pg.extract_text() or "").strip()
            except Exception:
                t = ""
            if t:
                partes.append(f"\n── página {i + 1} ──\n{t}")
        if partes:
            return titulo + "\n".join(partes)
    except ImportError:
        pass
    except Exception as exc:
        log.warning("pypdf falhou (%s), tentando pypdfium2", exc)
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(io.BytesIO(raw))
        for i in range(min(len(doc), 200)):
            try:
                t = (doc[i].get_textpage().get_text_range() or "").strip()
            except Exception:
                t = ""
            if t:
                partes.append(f"\n── página {i + 1} ──\n{t}")
        if partes:
            return titulo + "\n".join(partes)
    except ImportError:
        pass
    if partes:
        return "\n".join(partes)
    raise DocErro("Não consegui ler este PDF (pode ser só imagem escaneada). "
                  "Falta instalar? `pip install pypdf pypdfium2`.", "pdf_ilegivel")


def _extrair_docx(raw: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise DocErro("Para DOCX instale: `pip install python-docx`.", "sem_lib") from None
    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception:
        raise DocErro("DOCX inválido ou corrompido.", "invalido") from None
    blocos = []
    title = getattr(doc.core_properties, "title", None)
    if title and str(title).strip():
        blocos.append(f"# {str(title).strip()[:200]}")
    for p in doc.paragraphs:
        if not p.text or not p.text.strip():
            continue
        style = getattr(getattr(p, "style", None), "name", "") or ""
        if style.lower().startswith("heading"):
            level = re.search(r"\d+", style)
            blocos.append(f"{'#' * min(int(level.group()) if level else 1, 6)} {p.text.strip()}")
        else:
            blocos.append(p.text.strip())
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocos.append("| " + " | ".join(cells) + " |")
    texto = "\n".join(blocos).strip()
    if not texto:
        raise DocErro("DOCX sem texto legível.", "vazio")
    return texto


def _extrair_xlsx(raw: bytes, _nome: str) -> str:
    del _nome
    try:
        import openpyxl
    except ImportError:
        raise DocErro("Para planilhas instale: `pip install openpyxl pandas`.", "sem_lib") from None
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise DocErro("Planilha inválida ou corrompida.", "invalido") from None
    out = []
    title = getattr(wb.properties, "title", None)
    if title and str(title).strip():
        out.append(f"# {str(title).strip()[:200]}")
    for ws in wb.worksheets[:10]:
        out.append(f"\n## {ws.title}\n")
        n = 0
        for row in ws.iter_rows(values_only=True):
            vals = [(str(v).strip() if v is not None else "") for v in list(row)[:20]]
            if any(vals):
                out.append(" | ".join(vals))
                n += 1
                if n >= 200:
                    out.append("… (linhas restantes omitidas)")
                    break
    texto = "\n".join(out).strip()
    if not texto:
        raise DocErro("Planilha sem dados legíveis.", "vazio")
    return texto


def _extrair_csv(raw: bytes, delim: str = ",") -> str:
    texto = _decode(raw)
    try:
        linhas = list(csv.reader(io.StringIO(texto), delimiter=delim))[:201]
    except Exception:
        raise DocErro("CSV inválido.", "invalido") from None
    if not linhas:
        raise DocErro("CSV vazio.", "vazio")
    out = [" | ".join(c.strip() for c in r[:20]) for r in linhas if any((c or "").strip() for c in r)]
    if len(linhas) == 201:
        out.append("… (linhas restantes omitidas)")
    return "\n".join(out)


def _extrair_html(raw: bytes) -> str:
    texto = _decode(raw)
    try:
        from bs4 import BeautifulSoup
        sopa = BeautifulSoup(texto, "html.parser")
        for tag in sopa(["script", "style", "nav", "footer", "header", "iframe",
                         "object", "embed", "form", "template", "svg", "canvas"]):
            tag.decompose()
        for tag in sopa.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
            tag.insert_before("\n" + "#" * int(tag.name[1]) + " ")
        title = sopa.title.get_text(" ", strip=True) if sopa.title else ""
        body = re.sub(r"\n{3,}", "\n\n", sopa.get_text("\n", strip=True))
        return (f"# {title}\n\n" if title else "") + body
    except ImportError:
        # script/style sem fechar valem até EOF (raw text); demais tags exigem
        # fechamento para não devorar o documento em HTML malformado.
        txt = re.sub(r"(?is)<(script|style)\b[^>]*>(?:.*?</\1\s*>|.*$)", " ", texto)
        txt = re.sub(r"(?is)<(nav|footer|header|iframe|object|embed|form|template|svg|canvas)\b[^>]*>.*?</\1\s*>", " ", txt)
        title = re.search(r"(?is)<title\b[^>]*>(.*?)</title\s*>", txt)
        headings = re.sub(r"(?is)<h([1-6])\b[^>]*>(.*?)</h\1\s*>",
                          lambda m: "\n" + "#" * int(m.group(1)) + " " + re.sub(r"<[^>]+>", " ", m.group(2)), txt)
        body = re.sub(r"<[^>]+>", " ", headings)
        return (f"# {html.unescape(title.group(1)).strip()}\n\n" if title else "") + re.sub(
            r"\s+", " ", html.unescape(body)).strip()


def _extrair_json(raw: bytes) -> str:
    def rejeitar_constante(valor):
        raise ValueError(f"Constante JSON inválida: {valor}")

    try:
        objeto = json.loads(_decode(raw), parse_constant=rejeitar_constante)
    except (ValueError, TypeError, RecursionError):
        raise DocErro("JSON inválido.", "invalido") from None
    linhas = []

    def percorrer(valor, caminho=""):
        if len(linhas) >= 20000:
            return
        if isinstance(valor, dict):
            for chave, item in valor.items():
                nome = str(chave)[:200]
                if nome.lower() == "title" and isinstance(item, (str, int, float)):
                    linhas.append(f"# {item}")
                percorrer(item, f"{caminho}.{nome}" if caminho else nome)
        elif isinstance(valor, list):
            for indice, item in enumerate(valor[:5000]):
                percorrer(item, f"{caminho}[{indice}]")
        elif valor is not None:
            conteudo = str(valor).replace("\r", " ").strip()
            if conteudo:
                linhas.append(f"{caminho}: {conteudo}" if caminho else conteudo)

    percorrer(objeto)
    return "\n".join(linhas)


def extrair_texto(ext: str, raw: bytes, nome: str) -> str:
    if ext == ".pdf":
        texto = _extrair_pdf(raw)
    elif ext == ".docx":
        texto = _extrair_docx(raw)
    elif ext in (".xlsx",):
        texto = _extrair_xlsx(raw, nome)
    elif ext == ".xls":
        raise DocErro("Converta .xls para .xlsx ou .csv e envie de novo.", "formato_antigo")
    elif ext in (".csv", ".tsv"):
        texto = _extrair_csv(raw, "\t" if ext == ".tsv" else ",")
    elif ext in (".html", ".htm"):
        texto = _extrair_html(raw)
    elif ext == ".json":
        texto = _extrair_json(raw)
    elif ext in (".txt", ".md", ".markdown"):
        texto = _decode(raw).strip()
        if not texto:
            raise DocErro("Arquivo de texto vazio.", "vazio")
    elif ext in (".pptx", ".ppt", ".odt", ".ods", ".epub"):
        raise DocErro(f"{ext} ainda não é lido. Converta para PDF e envie de novo.", "nao_suportado")
    else:
        raise DocErro(f"Tipo {ext or '?'} não suportado. Envie PDF, DOCX, TXT, CSV, XLSX, HTML ou JSON.",
                      "nao_suportado")
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    if len(texto) < 10:
        raise DocErro("Não encontrei texto legível neste arquivo.", "vazio")
    return texto[:MAX_TEXTO_DOC]


def dividir_chunks(texto: str) -> list:
    partes, i, n = [], 0, len(texto)
    while i < n:
        partes.append(texto[i:i + TAM_CHUNK])
        if i + TAM_CHUNK >= n:
            break
        i += TAM_CHUNK - OVERLAP
    return partes


def _extrair_pagina_marcador(linha: str):
    m = re.fullmatch(r"\s*[─—-]{2,}\s*página\s+(\d+)\s*[─—-]{2,}\s*", linha, re.I)
    return int(m.group(1)) if m else None


def _chunk_records(texto: str, doc_id: str, source: str, version: int = DOC_VERSION) -> list:
    """Separa por seção/parágrafo e mantém página/overlap ao dividir regiões longas."""
    records = []
    section = ""
    page = None
    region = []
    region_section = ""
    region_page = None

    def emit_region():
        nonlocal region
        if not region:
            return
        blocos = []
        atual = ""
        limite_body = TAM_CHUNK - OVERLAP
        for paragraph in region:
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            if atual and len(atual) + len(paragraph) + 1 > limite_body:
                blocos.append(atual)
                atual = ""
            if len(paragraph) > limite_body:
                if atual:
                    blocos.append(atual)
                    atual = ""
                remaining = paragraph
                while len(remaining) > limite_body:
                    cut = remaining.rfind(" ", 0, limite_body + 1)
                    if cut < max(1, limite_body // 2):
                        cut = limite_body
                    blocos.append(remaining[:cut].strip())
                    remaining = remaining[cut:].strip()
                if remaining:
                    atual = remaining
            else:
                atual = f"{atual}\n{paragraph}".strip()
        if atual:
            blocos.append(atual)
        prefixo = (region_section.strip() + "\n") if region_section.strip() else ""
        prev = ""
        for pos, body in enumerate(blocos):
            if pos == 0 and prefixo and not body.startswith(region_section.strip()):
                body = (prefixo + body)[:TAM_CHUNK]
            overlap = min(OVERLAP, len(prev)) if prev else 0
            full = (prev[-overlap:] if overlap else "") + body
            if len(full) > TAM_CHUNK:
                full = full[:TAM_CHUNK]
            idx = len(records)
            digest = hashlib.sha256(full.encode("utf-8")).hexdigest()
            records.append({
                "idx": idx, "chunk_id": f"{doc_id}:{idx}", "text": full,
                "section": region_section, "page": region_page,
                "content_hash": digest, "version": version, "source": source,
                "table_hint": int("|" in body or bool(re.search(r"\b(tabela|table|sheet)\b", region_section, re.I))),
                "overlap": overlap,
            })
            prev = body
        region = []

    for line in texto.splitlines():
        page_marker = _extrair_pagina_marcador(line)
        heading = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if page_marker is not None or heading:
            if region and ((page_marker is not None and page_marker != region_page)
                           or (heading and heading.group(2).strip() != region_section)):
                emit_region()
            if page_marker is not None:
                page = page_marker
            if heading:
                section = heading.group(2).strip()[:200]
            if not region:
                region_section, region_page = section, page
            continue
        if not line.strip():
            continue
        if not region:
            region_section, region_page = section, page
        region.append(line)
    emit_region()
    return records


# ── Persistência ──────────────────────────────────────────────────────
def salvar_documento(db, nome: str, raw: bytes, conversation_id: str, owner=DONO_PADRAO) -> dict:
    import uuid
    owner = (str(owner or "").strip().lower() or DONO_PADRAO)[:40]
    if not isinstance(raw, (bytes, bytearray)) or not raw:
        raise DocErro("Arquivo vazio.", "vazio")
    if len(raw) > MAX_BYTES:
        raise DocErro(f"Arquivo grande demais ({len(raw) // 1048576}MB; limite 15MB).", "grande")
    base = re.sub(r'[\\/:*?"<>|]', "_", (nome or "arquivo").strip())[:120] or "arquivo"
    ext = ("." + base.rsplit(".", 1)[1].lower()) if "." in base else ""
    if ext not in EXT_OK:
        raise DocErro(f"Tipo {ext or '?'} não suportado. Envie PDF, DOCX, TXT, CSV, XLSX, HTML ou JSON.",
                      "nao_suportado")
    _validar_conteudo(ext, bytes(raw))
    texto_completo = extrair_texto(ext, bytes(raw), base)
    truncado = len(texto_completo) >= MAX_TEXTO_DOC
    texto = texto_completo[:MAX_TEXTO_DOC]
    doc_id = "doc-" + uuid.uuid4().hex[:12]
    records = _chunk_records(texto, doc_id, base)
    headings = [{"section": r["section"], "page": r["page"]}
                for r in records if r["section"]]
    metadata = {
        "sections": list(dict.fromkeys(
            (h["section"], h["page"]) for h in headings
        )),
        "pages": max((r["page"] or 0 for r in records), default=0),
        "table_hints": [r["idx"] for r in records if r["table_hint"]],
        "truncated": truncado,
        "extraction": "text_only",
    }
    metadata["sections"] = [{"section": title, "page": pg} for title, pg in metadata["sections"]]
    title = next((r["section"] for r in records if r["section"]), Path(base).stem[:200])
    content_hash = hashlib.sha256(texto.encode("utf-8")).hexdigest()
    try:
        pdf_title = re.search(r"(?im)^\s*title\s*[:=]\s*(.+)$", texto[:1000])
        if ext == ".json":
            parsed = json.loads(_decode(bytes(raw)))
            if isinstance(parsed, dict) and isinstance(parsed.get("title"), str):
                title = parsed["title"][:200]
        elif pdf_title:
            title = pdf_title.group(1).strip()[:200]
    except (ValueError, TypeError):
        pass
    with db:
        db.execute("""INSERT INTO documents
            (id, name, ext, size, chars, chunks, conversation_id, created_at, owner_id,
             title, extracted_meta, content_hash, version, source)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (doc_id, base, ext, len(raw), len(texto), len(records), conversation_id,
             agora_iso(), owner, title, json.dumps(metadata, ensure_ascii=False),
             content_hash, DOC_VERSION, base))
        for record in records:
            db.execute("""INSERT INTO doc_chunks
                (doc_id, idx, text, owner_id, chunk_id, section, page, content_hash,
                 version, source, table_hint, overlap)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (doc_id, record["idx"], record["text"], owner, record["chunk_id"],
                 record["section"], record["page"], record["content_hash"],
                 record["version"], record["source"], record["table_hint"], record["overlap"]))
            try:
                db.execute("INSERT INTO doc_chunks_fts(chunk_id, doc_id, owner_id, idx, text) "
                           "VALUES (?, ?, ?, ?, ?)",
                           (record["chunk_id"], doc_id, owner, record["idx"], record["text"]))
            except Exception:
                pass
    try:
        pasta = docs_dir()
        pasta.mkdir(parents=True, exist_ok=True)
        (pasta / (doc_id + ext)).write_bytes(bytes(raw))
    except OSError:
        pass
    log.info("documento %s (%s, %d chars, %d chunks)", doc_id, base, len(texto), len(records))
    return {"doc_id": doc_id, "name": base, "title": title, "ext": ext, "chars": len(texto),
            "chunks": len(records), "paginas_aprox": metadata["pages"] or max(1, len(texto) // 3000),
            "sections": metadata["sections"], "table_hints": len(metadata["table_hints"]),
            "content_hash": content_hash, "version": DOC_VERSION}


def listar_documentos(db, limite: int = 50, owner=DONO_PADRAO) -> list:
    rows = db.execute("""SELECT id, name, ext, size, chars, chunks,
                                summary IS NOT NULL, conversation_id, created_at, title,
                                content_hash, version
                         FROM documents WHERE owner_id=? ORDER BY created_at DESC LIMIT ?""",
                      (owner, min(max(int(limite), 1), 200))).fetchall()
    return [{"doc_id": i, "name": n, "ext": e, "size": s, "chars": c,
             "chunks": k, "tem_resumo": bool(r), "conversation_id": cv,
             "created_at": ca, "title": title, "content_hash": digest, "version": version}
            for i, n, e, s, c, k, r, cv, ca, title, digest, version in rows]


def obter_texto(db, doc_id: str, limite: int = 60000, owner=DONO_PADRAO) -> str:
    rows = db.execute("SELECT text, overlap FROM doc_chunks WHERE doc_id=? AND owner_id=? ORDER BY idx",
                      (doc_id, owner)).fetchall()
    partes = []
    for texto, overlap in rows:
        if overlap:
            prefix = texto[:overlap]
            if partes and partes[-1].endswith(prefix):
                texto = texto[overlap:]
            else:
                # Rebuilds from pre-V4.8 chunks have no stored overlap metadata.
                for n in range(min(OVERLAP, len(partes[-1]) if partes else 0), 0, -1):
                    if partes and partes[-1].endswith(texto[:n]):
                        texto = texto[n:]
                        break
        partes.append(texto)
    return "".join(partes)[:limite]


def trechos_relevantes(db, pergunta: str, doc_id: str = "", k: int = 4, owner=DONO_PADRAO, contexto: str = "") -> list:
    """Recupera trechos lexicais com proveniência, usando FTS5 quando disponível."""
    temporal = bool(re.search(r"\b(atual|atuais|recentes?|últim[oa]s?|latest|recent|current)\b", pergunta + " " + contexto, re.I))
    toks_q = _tokens(pergunta)
    if temporal:
        toks_q -= {"atual", "atuais", "recente", "recentes", "último", "últimos", "última", "últimas", "latest", "recent", "current"}
    if not toks_q:
        return []
    k = min(max(int(k), 1), 20)
    rows = []
    fts_available = False
    terms = sorted(toks_q)
    fts_query = (" OR " if temporal else " AND ").join('"' + term.replace('"', '""') + '"' for term in terms)
    try:
        fts_sql = """SELECT c.doc_id, c.idx, c.text, c.chunk_id, c.section, c.page,
                            c.content_hash, c.version, c.source, d.name, c.table_hint,
                            bm25(doc_chunks_fts)
                     FROM doc_chunks_fts
                     JOIN doc_chunks c ON c.chunk_id=doc_chunks_fts.chunk_id
                                       AND c.doc_id=doc_chunks_fts.doc_id
                                       AND c.owner_id=doc_chunks_fts.owner_id
                                       AND c.idx=doc_chunks_fts.idx
                     JOIN documents d ON d.id=c.doc_id AND d.owner_id=c.owner_id
                     WHERE doc_chunks_fts MATCH ? AND c.owner_id=?"""
        args = [fts_query, owner]
        if doc_id:
            fts_sql += " AND c.doc_id=?"
            args.append(doc_id)
        fts_sql += " ORDER BY bm25(doc_chunks_fts), c.idx LIMIT ?"
        rows = db.execute(fts_sql, (*args, min(100, k * 5) if temporal else k)).fetchall()
        fts_available = True
    except Exception:
        pass
    scored = []
    if rows:
        for row in rows:
            did, idx, txt, chunk_id, section, page, digest, version, source, name, table_hint, _ = row
            scored.append((0, did, idx, txt, chunk_id, section, page, digest, version, source, name, table_hint))
    else:
        sql = """SELECT c.doc_id, c.idx, c.text, c.chunk_id, c.section, c.page,
                        c.content_hash, c.version, c.source, d.name, c.table_hint
                 FROM doc_chunks c JOIN documents d ON d.id=c.doc_id AND d.owner_id=c.owner_id
                 WHERE c.owner_id=?"""
        args = [owner]
        if doc_id:
            sql += " AND c.doc_id=?"
            args.append(doc_id)
        fallback_rows = db.execute(sql, args).fetchall()
        for did, idx, txt, chunk_id, section, page, digest, version, source, name, table_hint in fallback_rows:
            toks = _tokens(txt)
            inter = toks_q & toks
            if inter:
                score = len(inter) / max(1, len(toks_q)) + 0.1 * len(inter) / max(1, len(toks))
                scored.append((score, did, idx, txt, chunk_id, section, page, digest, version, source, name,
                               table_hint))
    if not rows:
        scored.sort(key=lambda x: (-x[0], x[1], x[2]))
    if temporal:
        def temporal_key(row):
            # Explicit dates in content, never upload/index timestamps. Retain relevance
            # ahead of freshness; a recent unrelated chunk cannot win.
            text = row[3]
            coverage = len(toks_q & _tokens(text)) / max(1, len(toks_q))
            match = re.search(r"\b(?:report|relatório|published|publicado|updated|atualizado)\s*(?:em|in|:)?\s*((?:19|20)\d{2})(?:-\d{2}-\d{2})?", text[:500], re.I)
            from datetime import datetime, timezone
            year = int(match.group(1)) if match else 0
            if year > datetime.now(timezone.utc).year:
                year = 0
            return (-int(coverage >= 0.6), -year, -coverage, row[1], row[2])
        scored.sort(key=temporal_key)
    return [{"doc_id": did, "doc": name, "idx": idx, "chunk_id": chunk_id,
             "owner_id": owner, "section": section, "page": page, "content_hash": digest,
             "version": version, "source": source, "table_hint": bool(table_hint),
             "trecho": txt[:1500], "citacao": f"[{chunk_id}]"}
            for _, did, idx, txt, chunk_id, section, page, digest, version, source, name, table_hint in scored[:k]]


def comparar_documentos(db, pergunta: str, doc_ids: list, owner=DONO_PADRAO) -> dict:
    """Compara evidências lexicais sem inferir ausência semântica."""
    ids = list(dict.fromkeys(str(i)[:64] for i in (doc_ids or []) if i))
    if not ids or len(ids) > 20:
        raise DocErro("Informe de 1 a 20 documentos.", "limite")
    findings = []
    for doc_id in ids:
        doc = db.execute("""SELECT name, title, extracted_meta, content_hash, version
                             FROM documents WHERE id=? AND owner_id=?""", (doc_id, owner)).fetchone()
        if not doc:
            # Nunca revelar se o documento existe para outro dono.
            continue
        name, title, raw_meta, digest, version = doc
        meta = json.loads(raw_meta or "{}")
        evidence = trechos_relevantes(db, pergunta, doc_id, k=3, owner=owner)
        content = obter_texto(db, doc_id, MAX_TEXTO_DOC + 1, owner)
        literal = re.sub(r"\s+", " ", (pergunta or "").strip().strip('"“”')).casefold()
        normalized = re.sub(r"\s+", " ", content).casefold()
        if evidence:
            status = "located"
        elif literal and len(literal) <= 160 and not meta.get("truncated") and literal not in normalized:
            status = "proven_absent"
        else:
            status = "not_located"
        findings.append({
            "doc_id": doc_id, "name": name, "title": title, "status": status,
            "basis": "lexical_chunks" if evidence else (
                "exact_literal_not_present_in_complete_extraction" if status == "proven_absent"
                else "no_matching_evidence_or_incomplete_extraction"),
            "content_hash": digest, "version": version, "evidence": evidence,
        })
    return {"query": pergunta[:500], "owner_id": owner, "findings": findings,
            "note": "Busca lexical; proven_absent significa apenas que a expressão literal não aparece no texto extraído completo, não ausência semântica."}


def excluir_documento(db, doc_id: str, owner=DONO_PADRAO) -> bool:
    """Remove documento e seus índices somente do proprietário autenticado."""
    row = db.execute("SELECT ext FROM documents WHERE id=? AND owner_id=?", (doc_id, owner)).fetchone()
    if not row:
        return False
    ext = row[0]
    with db:
        try:
            db.execute("DELETE FROM doc_chunks_fts WHERE doc_id=? AND owner_id=?", (doc_id, owner))
        except Exception:
            pass
        db.execute("DELETE FROM doc_chunks WHERE doc_id=? AND owner_id=?", (doc_id, owner))
        db.execute("DELETE FROM documents WHERE id=? AND owner_id=?", (doc_id, owner))
    try:
        (docs_dir() / (doc_id + ext)).unlink(missing_ok=True)
    except OSError:
        log.warning("não foi possível remover arquivo original do documento %s", doc_id)
    return True


# ── Resumo e tópicos via Hermes ───────────────────────────────────────
def _llm(cfg: dict, sistema: str, usuario: str, max_tokens: int = 800) -> str:
    import hermes_bridge
    res = hermes_bridge.chat_complete(
        cfg["proxy"], cfg["model"],
        [{"role": "system", "content": sistema}, {"role": "user", "content": usuario}],
        max_tokens=max_tokens)
    return res["texto"]


def resumir_documento(db, cfg: dict, doc_id: str, owner=DONO_PADRAO) -> str:
    row = db.execute("SELECT summary, name FROM documents WHERE id=? AND owner_id=?",
                     (doc_id, owner)).fetchone()
    if not row:
        raise DocErro("Documento não encontrado.", "nao_encontrado")
    if row[0]:
        return row[0]
    texto = obter_texto(db, doc_id, owner=owner)
    if not texto.strip():
        raise DocErro("Documento sem texto.", "vazio")
    chunks = dividir_chunks(texto)
    if len(chunks) > 6:
        # amostra distribuída (começo, meio, fim) para caber no tempo/quota
        idxs = sorted({0, 1, len(chunks) // 3, len(chunks) // 2,
                       2 * len(chunks) // 3, len(chunks) - 1})
        chunks = [chunks[i] for i in idxs if 0 <= i < len(chunks)]
    parciais = []
    for ch in chunks:
        try:
            parciais.append(_llm(cfg, "Resuma o trecho em até 8 linhas, só o essencial.",
                                 ch[:4000], max_tokens=500))
        except Exception as exc:
            log.warning("resumo parcial falhou: %s", exc)
    if not parciais:
        raise DocErro("A IA falhou ao resumir. Tente de novo.", "ia_falhou")
    if len(parciais) == 1:
        final = parciais[0]
    else:
        final = _llm(cfg, "Junte os resumos parciais num resumo único, organizado com títulos e bullets.",
                      "\n\n---\n\n".join(parciais)[:15000], max_tokens=1000)
    with db:
        db.execute("UPDATE documents SET summary=? WHERE id=? AND owner_id=?", (final, doc_id, owner))
    return final


def topicos_documento(db, cfg: dict, doc_id: str, owner=DONO_PADRAO) -> str:
    row = db.execute("SELECT topics, name FROM documents WHERE id=? AND owner_id=?",
                     (doc_id, owner)).fetchone()
    if not row:
        raise DocErro("Documento não encontrado.", "nao_encontrado")
    if row[0]:
        return row[0]
    texto = obter_texto(db, doc_id, 30000, owner)
    try:
        final = _llm(cfg, "Extraia do documento: 1) tópicos principais (bullets), "
                           "2) palavras-chave, 3) datas/números importantes, "
                           "4) pendências ou próximos passos, se houver. Markdown curto.",
                      texto, max_tokens=800)
    except Exception as exc:
        raise DocErro(f"A IA falhou: {exc}.", "ia_falhou") from None
    with db:
        db.execute("UPDATE documents SET topics=? WHERE id=? AND owner_id=?", (final, doc_id, owner))
    return final
